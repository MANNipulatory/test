import streamlit as st
import os
import io
import re
import time
from google import genai
from google.genai import types
from google.genai.errors import APIError as GeminiAPIError
from openai import OpenAI
import pypdf
from docx import Document  # 📄 เพิ่ม Library สำหรับเสกไฟล์ Word (.docx)

# ── 1. INITIALIZATION (HYBRID CREW) ──────────────────────────────
if "GEMINI_API_KEY" in st.secrets and st.secrets["GEMINI_API_KEY"]:
    gemini_client = genai.Client(api_key=st.secrets["GEMINI_API_KEY"])
elif os.environ.get("GEMINI_API_KEY"):
    gemini_client = genai.Client(api_key=os.environ.get("GEMINI_API_KEY"))
else:
    gemini_client = None

GEMINI_MODEL = "gemini-2.5-flash"

if "TYPHOON_API_KEY" in st.secrets and st.secrets["TYPHOON_API_KEY"]:
    typhoon_key = st.secrets["TYPHOON_API_KEY"]
elif os.environ.get("TYPHOON_API_KEY"):
    typhoon_key = os.environ.get("TYPHOON_API_KEY")
else:
    typhoon_key = None

if typhoon_key:
    typhoon_client = OpenAI(api_key=typhoon_key, base_url="https://api.opentyphoon.ai/v1")
else:
    typhoon_client = None

TYPHOON_MODEL = "typhoon-v1.5x-70b-instruct"

st.set_page_config(
    page_title="AI Procurement TOR Workspace",
    page_icon="🛡️",
    layout="wide"
)

# 🎨 [ELEVATED MINIMALISM DESIGN]
st.markdown("""
    <style>
        .block-container { max-width: 1000px !important; padding-top: 2.5rem !important; padding-bottom: 5rem !important; }
        h2 { font-size: 1.3rem !important; font-weight: 700 !important; color: #1E3A8A; margin-top: 2.5rem !important; margin-bottom: 1.2rem !important; padding-left: 12px !important; border-left: 4px solid #3B82F6 !important; }
        .stTextArea textarea { background-color: rgba(128, 128, 128, 0.06) !important; border: 1px solid rgba(128, 128, 128, 0.2) !important; border-radius: 8px !important; padding: 12px !important; font-size: 0.95rem !important; }
        .stTextArea textarea:focus { border-color: #3B82F6 !important; background-color: transparent !important; box-shadow: 0 0 0 3px rgba(59, 130, 246, 0.15) !important; }
        .stTextInput input { background-color: rgba(128, 128, 128, 0.06) !important; border: 1px solid rgba(128, 128, 128, 0.2) !important; border-radius: 8px !important; }
        .streamlit-expanderHeader { background-color: rgba(128, 128, 128, 0.04) !important; border: 1px solid rgba(128, 128, 128, 0.15) !important; border-radius: 8px !important; font-weight: 600 !important; }
        .stButton button { border-radius: 6px !important; font-size: 0.9rem !important; font-weight: 500 !important; }
        .stButton div button[data-testid="baseButton-primary"] { background: linear-gradient(135deg, #2563EB, #1D4ED8) !important; color: #FFFFFF !important; border: none !important; box-shadow: 0 4px 6px rgba(37, 99, 235, 0.2) !important; }
        
        .export-box {
            background-color: rgba(59, 130, 246, 0.04) !important;
            border: 1px dashed rgba(59, 130, 246, 0.3) !important;
            border-radius: 12px !important;
            padding: 20px !important;
            margin-top: 30px;
        }
    </style>
""", unsafe_allow_html=True)

TOR_TITLES = {
    1: "ความเป็นมา", 2: "วัตถุประสงค์", 3: "คุณสมบัติผู้ยื่นข้อเสนอ",
    4: "แบบรูปรายการ หรือคุณลักษณะเฉพาะของพัสดุ", 5: "ระยะเวลาดำเนินการ",
    6: "ระยะเวลาส่งมอบงาน หรือส่งมอบพัสดุ", 7: "หลักเกณฑ์ในการพิจารณาคัดเลือกข้อเสนอ",
    8: "อัตราค่าปรับ", 9: "การรับประกันความชำรุดบกพร่อง", 10: "ข้อสงวนสิทธิ์ในการยื่นข้อเสนอและอื่นๆ"
}

SYSTEM_PROMPT = (
    "คุณคือผู้เชี่ยวชาญระดับสูงด้านการจัดซื้อจัดจ้างภาครัฐไทย "
    "หน้าที่ของคุณคือเขียนเนื้อหาขอบเขตของงาน (TOR) ตามมาตรฐานหนังสือเวียน กค (กวจ) 0405.4/ว 159 ของกรมบัญชีกลาง "
    "จงเขียนอธิบายอย่างละเอียด ถี่ถ้วน ครอบคลุมทุกมิติทางกฎหมายพัสดุ ใช้ภาษาราชการไทยที่เป็นทางการอย่างสมบูรณ์ "
    "ห้ามระบุยี่ห้อหรือรุ่นของสินค้าโดยตรงเด็ดขาด ให้ใช้คุณลักษณะเชิงฟังก์ชัน (Functional Specification) เสมอ "
    "ใช้ตัวเลขอารบิกในการรันข้อย่อย ห้ามใช้เลขไทยในทุกกรณี "
    "ห้ามแปลเนื้อหาเป็นภาษาอังกฤษล้วน ให้ตอบเป็นภาษาไทยอย่างเป็นทางการเท่านั้น"
)

# ── 2. STATE MANAGEMENT ──────────────────────────────────────────
if "tor_sections" not in st.session_state: st.session_state.tor_sections = {i: "" for i in range(1, 11)}
if "meta_data" not in st.session_state: st.session_state.meta_data = {}
if "pdf_extracted_texts" not in st.session_state: st.session_state.pdf_extracted_texts = {}
if "ai_drafted_spec_v4" not in st.session_state: st.session_state.ai_drafted_spec_v4 = ""

# ── 3. HELPERS & GENERATORS ──────────────────────────────────────
def call_gemini_with_retry(func, *args, **kwargs):
    max_retries = 3
    for attempt in range(max_retries):
        try: return func(*args, **kwargs)
        except GeminiAPIError as e:
            if e.code in [429, 503]: time.sleep(5); continue
            raise e

def call_typhoon_with_retry(func, *args, **kwargs):
    max_retries = 5
    for attempt in range(max_retries):
        try: return func(*args, **kwargs)
        except Exception as e:
            err_msg = str(e).lower()
            if any(k in err_msg for k in ["rate_limit", "429", "503", "overloaded"]):
                time.sleep(15); continue
            raise e

def extract_text_from_pdf(uploaded_file):
    try:
        pdf_reader = pypdf.PdfReader(uploaded_file)
        return "".join([page.extract_text() or "" for page in pdf_reader.pages])
    except Exception as e: return ""

# 🛡️ ฟังก์ชันกรองความโปร่งใสตัวเดิม คืนชีพกลับมาทำภารกิจแบบรัดกุม 100%
def clean_redundant_info(text):
    if not text: return ""
    # 1. กรองหมายเลขโทรศัพท์รูปแบบต่างๆ
    text = re.sub(r'\b\d{2,3}-\d{3}-\d{4}\b|\b\d{2,3}-\d{4}-\d{4}\b|\b\d{9,10}\b', "[PHONE_HIDDEN]", text)
    # 2. กรองที่อยู่อีเมล
    text = re.sub(r'[\w\.-]+@[\w\.-]+\.\w+', "[EMAIL_HIDDEN]", text)
    # 3. กรองลิงก์และ URL
    text = re.sub(r'https?://[^\s<>"]+|www\.[^\s<>"]+', "[URL_HIDDEN]", text)
    # 4. กรองชื่อบริษัทภาษาไทย หจก. บมจ. บจก. ทุกรูปแบบโครงสร้าง
    text = re.sub(r'(บริษัท\s+[^\s\n]+?\s+จำกัด(?:\s*\(มหาชน\))?)|(ห้างหุ้นส่วนจำกัด\s+[^\s\n]+)|(\bหจก\s*\.\s*[^\s\n]+)|(\bบมจ\s*\.\s*[^\s\n]+)|(\bบจก\s*\.\s*[^\s\n]+)', "[COMPANY_HIDDEN]", text)
    # 5. กรองชื่อบริษัทภาษาอังกฤษ (Co., Ltd. / Inc. / Corp.)
    en_company = r'\b[A-Za-z0-9\s\.,&\-\(\)]+?\s+(?:Co\s*\.?\s*,?\s*Ltd\s*\.?|Company\s+Limited|Inc\s*\.?|Corp\s*\.?|LLC|Group)\b'
    text = re.sub(en_company, "[COMPANY_HIDDEN]", text, flags=re.IGNORECASE)
    # 6. กรองรายชื่อแบรนด์และเครื่องหมายการค้าชั้นนำ เพื่อป้องกันการล็อกสเปคตามกฎพัสดุ ว.159
    text = re.sub(r'\b(Intel|AMD|NVIDIA|GeForce|Asus|Acer|HP|Dell|Lenovo|Apple|Microsoft|Cisco|Huawei)\b', "[BRAND_HIDDEN]", text, flags=re.IGNORECASE)
    return text

def summarize_with_gemini(raw_text):
    if not gemini_client or not raw_text: return raw_text
    prompt = f"""
    คุณคือผู้ช่วยสกัดข้อมูลทางเทคนิค อ่านข้อความด้านล่างนี้แล้วสรุปเนื้อหาสำคัญเกี่ยวกับ 'รายละเอียดคุณลักษณะเฉพาะทางเทคนิคทั้งหมด' 
    และ 'เงื่อนไขการรับประกันและการสนับสนุน' โดยรักษาข้อมูลตัวเลขสเปคทางเทคนิคไว้ให้ครบถ้วนที่สุด แต่ตัดพวกเศษขยะหรือชื่อคู่ค้าออก
    [ข้อมูลเอกสาร]:
    {raw_text[:12000]}
    """
    try:
        response = call_gemini_with_retry(
            gemini_client.models.generate_content, model=GEMINI_MODEL, contents=prompt,
            config=types.GenerateContentConfig(temperature=0.1)
        )
        return response.text
    except Exception: return raw_text[:3000]

def generate_typhoon_stream(prompt_text, section_num, placeholder):
    if not typhoon_client: return ""
    full_response = ""
    try:
        response_stream = call_typhoon_with_retry(
            typhoon_client.chat.completions.create,
            model=TYPHOON_MODEL,
            messages=[{"role": "system", "content": SYSTEM_PROMPT}, {"role": "user", "content": prompt_text}],
            temperature=0.4, stream=True
        )
        for chunk in response_stream:
            if chunk.choices[0].delta.content:
                full_response += chunk.choices[0].delta.content
                placeholder.code(full_response + "▌", language="text")
        placeholder.code(full_response, language="text")
        st.session_state.tor_sections[section_num] = full_response
        return full_response
    except Exception as e: return ""

# 📄 [EXPORT GENERATORS] ฟังก์ชันแปลงข้อมูลสำหรับดาวน์โหลดเป็นฟอร์แมตต่างๆ
def build_word_document():
    doc = Document()
    title_p = doc.add_paragraph()
    title_run = title_p.add_run(f"ร่างขอบเขตของงาน (TOR)\nโครงการ: {st.session_state.meta_data.get('title', 'ไม่ได้ระบุ')}")
    title_run.bold = True
    title_run.font.size = 203200  # ประมาณ 16pt
    
    doc.add_paragraph(f"หน่วยงานเจ้าของโครงการ: {st.session_state.meta_data.get('agency', '-')}")
    doc.add_paragraph(f"วงเงินงบประมาณ: {st.session_state.meta_data.get('budget', 0):,} บาท")
    doc.add_paragraph("-" * 40)
    
    for idx in range(1, 11):
        heading = doc.add_paragraph()
        h_run = heading.add_run(f"ข้อ {idx} {TOR_TITLES[idx]}")
        h_run.bold = True
        
        content = st.session_state.tor_sections.get(idx, "ไม่มีเนื้อหา")
        doc.add_paragraph(content)
        doc.add_paragraph()
        
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

def build_pdf_html_document():
    html_content = f"""
    <html>
    <head>
        <meta charset="utf-8">
        <style>
            body {{ font-family: 'Sarabun', 'Tahama', sans-serif; padding: 40px; color: #333; line-height: 1.6; }}
            h1 {{ text-align: center; font-size: 22px; margin-bottom: 5px; }}
            .meta {{ text-align: center; font-size: 14px; color: #555; margin-bottom: 30px; }}
            .section {{ margin-bottom: 25px; page-break-inside: avoid; }}
            .title {{ font-weight: bold; font-size: 16px; color: #1E3A8A; border-bottom: 1px solid #ddd; padding-bottom: 5px; margin-bottom: 10px; }}
            .content {{ font-size: 14px; white-space: pre-wrap; text-align: justify; }}
        </style>
    </head>
    <body>
        <h1>ร่างขอบเขตของงาน (TOR)</h1>
        <div class="meta">
            <b>โครงการ:</b> {st.session_state.meta_data.get('title', '-')}<br>
            <b>หน่วยงาน:</b> {st.session_state.meta_data.get('agency', '-') or 'ไม่ได้ระบุ'} | 
            <b>งบประมาณ:</b> {st.session_state.meta_data.get('budget', 0):,} บาท
        </div>
    </body>
    </html>
    """
    for idx in range(1, 11):
        text_inside = st.session_state.tor_sections.get(idx, "").replace("\n", "<br>")
        html_content = html_content.replace("</body>", f"""
        <div class="section">
            <div class="title">ข้อ {idx} {TOR_TITLES[idx]}</div>
            <div class="content">{text_inside}</div>
        </div>
        </body>""")
    return html_content

# ── 4. UI WORKFLOW ───────────────────────────────────────────────
st.title("🛡️ AI Procurement TOR Workspace")
st.caption("⚡ Hybrid Core Mode: วิเคราะห์ด้วย Gemini ➔ ร่างข้อกำหนดภาษาราชการเนียนตาด้วย Typhoon")

# 📊 ส่วนที่ 1: ตั้งค่าข้อมูลโครงการหลัก
st.markdown("## 1. ข้อมูลโครงการทั่วไป")
col_form1, col_form2 = st.columns(2)
with col_form1:
    p_name = st.text_input("ชื่อโครงการ / งานจัดซื้อจัดจ้าง *", placeholder="ระบุชื่อโครงการ...")
    p_agency = st.text_input("หน่วยงาน / ส่วนราชการ", placeholder="ระบุหน่วยงานผู้รับผิดชอบ...")
with col_form2:
    p_budget = st.number_input("วงเงินงบประมาณ (บาท) *", min_value=0, step=5000, value=0)
    p_criteria = st.radio("หลักเกณฑ์การคัดเลือกข้อเสนอ", ["เกณฑ์ราคา", "เกณฑ์ราคาประกอบเกณฑ์อื่น"], horizontal=True)

# 🔍 ส่วนที่ 2: ใช้ Gemini สกัดเนื้อหา (ข้อ 4)
st.markdown("## 2. วิเคราะห์เอกสารสเปคกลางเพื่อสร้าง TOR ข้อ 4")
job_description_pdf = st.text_area("วัตถุประสงค์ / ลักษณะงานที่ต้องการใช้งานจริง:", placeholder="อธิบายเป้าหมายการใช้งานจริงเพื่อนำทาง AI...", height=80)

col_file1, col_file2 = st.columns([3, 1])
with col_file1:
    uploaded_files = st.file_uploader("อัปโหลดไฟล์ PDF อ้างอิงสเปค", type=["pdf"], accept_multiple_files=True, label_visibility="collapsed")
with col_file2:
    process_pdf_click = st.button("🔓 สกัดข้อมูลด้วย Gemini", use_container_width=True)

if uploaded_files and process_pdf_click:
    with st.spinner("🤖 Gemini กำลังวิเคราะห์และคลีนข้อมูลแบรนด์สินค้า..."):
        st.session_state.pdf_extracted_texts = {}
        for file in uploaded_files:
            raw_text = extract_text_from_pdf(file)
            cleaned_text = clean_redundant_info(raw_text)
            st.session_state.pdf_extracted_texts[file.name] = summarize_with_gemini(cleaned_text)

if st.session_state.pdf_extracted_texts:
    for file_name, text_content in list(st.session_state.pdf_extracted_texts.items()):
        with st.expander(f"🔎 ผลลัพธ์การสกัดสเปคอ้างอิง: {file_name}", expanded=False):
            user_updated_text = st.text_area("แก้ไขเนื้อหาเพิ่มเติม", value=text_content, height=120, key=f"edit_pdf_{file_name}", label_visibility="collapsed")
            st.session_state.pdf_extracted_texts[file_name] = user_updated_text
    
    check_company = st.checkbox("ข้าพเจ้ายืนยันว่าได้ตรวจสอบความโปร่งใสและตัดข้อมูลล็อกสเปคแล้ว")
    
    if st.button("📊 บันทึกสเปคกลางเข้าสู่ข้อ 4 TOR", type="primary"):
        if not job_description_pdf: st.warning("⚠️ โปรดระบุวัตถุประสงค์ลักษณะงานก่อน")
        elif not check_company: st.error("⚠️ โปรดกดยืนยันการตรวจสอบความโปร่งใสก่อน")
        else:
            with st.spinner("🤖 สรุปรวมข้อมูลเข้าสู่ข้อ 4 ด้วย Gemini..."):
                try:
                    all_data = "".join([f"\n[เอกสารอ้างอิง]\n{t}\n" for t in st.session_state.pdf_extracted_texts.values()])
                    prompt = f"วิเคราะห์และร่างสเปคกลางใส่ใน TOR ข้อ 4 วัตถุประสงค์คือ {job_description_pdf} ข้อมูลจากผู้เสนอคือ {all_data} เขียนแบ่งข้อย่อย 4.1, 4.2 อย่างละเอียดและห้ามมีชื่อแบรนด์เด็ดขาด"
                    response = call_gemini_with_retry(gemini_client.models.generate_content, model=GEMINI_MODEL, contents=prompt)
                    st.session_state.ai_drafted_spec_v4 = response.text
                    st.session_state.tor_sections[4] = response.text
                    st.success("ซิงค์สเปคเข้าสู่ร่างข้อ 4 เรียบร้อยแล้ว!")
                except Exception as e: st.error(f"เกิดข้อผิดพลาด: {e}")

# 📝 ส่วนที่ 3: ใช้ Typhoon ร่างเนื้อหา 10 ข้อหลัก
st.markdown("## 3. จัดการขอบเขตงาน TOR 10 ข้อหลัก")
if p_name: st.caption(f"📍 โครงการปัจจุบัน: **{p_name}** | งบประมาณ: **{p_budget:,} บาท**")

for i in range(1, 11):
    with st.expander(f"📌 ข้อ {i}: {TOR_TITLES[i]}", expanded=(i==4)):
        current_content = st.session_state.tor_sections.get(i, "")
        if i == 4 and st.session_state.ai_drafted_spec_v4:
            st.markdown("<small style='color:#10B981; font-weight:600;'>✓ ซิงค์ข้อมูลโครงสร้างสเปคทางเทคนิคเรียบร้อยแล้ว</small>", unsafe_allow_html=True)
            
        updated_content = st.text_area(f"เนื้อหาข้อ {i}", value=current_content, height=160, key=f"main_edit_sec_{i}", label_visibility="collapsed")
        st.session_state.tor_sections[i] = updated_content
        
        col_space, col_action = st.columns([5, 1.2])
        with col_action:
            if st.button(f"🔄 รีเจนเนื้อหาข้อ {i}", key=f"main_regen_btn_{i}", use_container_width=True):
                if not p_name: st.error("กรุณาระบุชื่อโครงการก่อน")
                else:
                    box_placeholder = st.empty()
                    spec_prompt = f"จงเขียนทบทวนร่างข้อกำหนดขอบเขตงาน TOR โครงการ '{p_name}' เฉพาะในส่วนของ 'ข้อ {i} หัวข้อ: {TOR_TITLES[i]}' ให้ออกมาเป็นข้อๆ ภาษาราชการเต็มรูปแบบ งบประมาณคือ {p_budget} บาท"
                    generate_typhoon_stream(spec_prompt, i, box_placeholder)
                    st.rerun()

st.write("")
if st.button("✨ ให้ Typhoon เริ่มร่างข้อกำหนดส่วนที่เหลือพร้อมกันทั้งหมด", type="primary", use_container_width=True):
    if not p_name: st.error("⚠️ กรุณาระบุชื่อโครงการที่ด้านบนสุดก่อนรันงาน")
    else:
        st.session_state.meta_data = {"title": p_name, "agency": p_agency, "type": "ซื้อ/จ้างทั่วไป", "budget": p_budget, "criteria": p_criteria}
        for i in range(1, 11):
            if i == 4 and st.session_state.tor_sections[4]: continue
            st.toast(f"Typhoon กำลังร่างเรียบเรียงข้อ {i}...")
            box_placeholder = st.empty()
            proj_context = f"โครงการ: {p_name}, หน่วยงาน: {p_agency}, งบประมาณ: {p_budget} บาท, เกณฑ์พิจารณา: {p_criteria}"
            specific_prompt = f"จงเขียนเนื้อหาของขอบเขตของงาน (TOR) ตามมาตรฐานราชการไทย ว.159 เฉพาะ 'ข้อ {i} หัวข้อ: {TOR_TITLES[i]}' ของ{proj_context} อธิบายรายละเอียดเชิงระเบียบพัสดุให้ครบถ้วนและสละสลวยที่สุด"
            generate_typhoon_stream(specific_prompt, i, box_placeholder)
            time.sleep(2.0)
        st.balloons()

# ── 5. EXPORT HUB (ศูนย์ดาวน์โหลด 3 รูปแบบ) ───────────────────────
if st.session_state.tor_sections[1] or st.session_state.tor_sections[4]:
    if not st.session_state.meta_data:
        st.session_state.meta_data = {"title": p_name or "ไม่ได้ระบุชื่อ", "agency": p_agency, "budget": p_budget, "criteria": p_criteria}
        
    st.markdown("<div class='export-box'>", unsafe_allow_html=True)
    st.markdown("### 📥 ศูนย์ส่งออกเอกสาร TOR (Export Hub)")
    st.caption("เลือกดาวน์โหลดไฟล์ขอบเขตงานในรูปแบบที่ต้องการไปใช้งานต่อ")
    
    all_text_export = f"ร่างขอบเขตของงาน (TOR) - {st.session_state.meta_data['title']}\n\n"
    for idx, ct in st.session_state.tor_sections.items():
        all_text_export += f"ข้อ {idx} {TOR_TITLES[idx]}\n{ct}\n\n"
        
    word_bytes = build_word_document()
    html_printable = build_pdf_html_document()
    
    dl_col1, dl_col2, dl_col3 = st.columns(3)
    
    with dl_col1:
        st.download_button(
            label="📝 ดาวน์โหลดไฟล์ Word (.DOCX)",
            data=word_bytes,
            file_name=f"TOR_{p_name or 'Project'}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True
        )
    with dl_col2:
        st.download_button(
            label="📄 ดาวน์โหลดไฟล์ข้อความ (.TXT)",
            data=all_text_export,
            file_name=f"TOR_{p_name or 'Project'}.txt",
            mime="text/plain",
            use_container_width=True
        )
    with dl_col3:
        st.download_button(
            label="🖨️ เปิดพิมพ์ / บันทึกเป็น PDF",
            data=html_printable,
            file_name=f"TOR_Print_Preview_{p_name or 'Project'}.html",
            mime="text/html",
            use_container_width=True,
            help="ดาวน์โหลดไฟล์พรีวิวนี้แล้วเปิดใน Chrome/Edge จากนั้นกด Ctrl+P เพื่อบันทึกเป็น PDF ที่ฟอนต์ไม่พังครับ"
        )
    st.markdown("</div>", unsafe_allow_html=True)
